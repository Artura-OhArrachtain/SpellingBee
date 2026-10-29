"""Web application for spelling and grammar assistance."""

from __future__ import annotations

import io
import json
import os
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from zipfile import ZIP_DEFLATED, ZipFile

from flask import Flask, jsonify, render_template, request, send_file, send_from_directory
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parents[1] / ".env")

app = Flask(__name__, template_folder="../templates", static_folder="../static")
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
MAX_TEXT_LENGTH = 12000
EXTENSION_DIR = Path(__file__).resolve().parents[2] / "integrations" / "browser-extension"


def extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension in {".txt", ".md", ".markdown"}:
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("This text file is not encoded as UTF-8.") from error

    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        sections = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    sections.append("\t".join(cells))
        return "\n".join(sections)

    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("This PDF is password-protected. Unlock it and try again.")
        if len(reader.pages) > 250:
            raise ValueError("This PDF has more than 250 pages. Please open a shorter document.")
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        if not text.strip():
            raise ValueError("No selectable text was found. Scanned image PDFs are not supported yet.")
        return text

    raise ValueError("Choose a TXT, Markdown, DOCX, or PDF document.")


def proofread_with_ai(text: str) -> dict[str, str]:
    """Ask an OpenAI-compatible chat completions endpoint to proofread text."""
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("AI proofreading is not configured. Set OPENAI_API_KEY to enable it.")

    endpoint = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    payload = {
        "model": model,
        "temperature": 0.2,
        "messages": [
            {
                "role": "system",
                "content": (
                    "Proofread the user's text for spelling, grammar, punctuation, and clarity. "
                    "Preserve its meaning, tone, and language. Return only valid JSON with "
                    "these string fields: corrected_text and explanation."
                ),
            },
            {"role": "user", "content": text},
        ],
        "response_format": {"type": "json_object"},
    }
    api_request = Request(
        f"{endpoint}/chat/completions",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urlopen(api_request, timeout=45) as response:
            result = json.loads(response.read().decode("utf-8"))
    except HTTPError as error:
        if error.code in (401, 403):
            raise RuntimeError("The AI service rejected the API key. Check OPENAI_API_KEY.") from error
        raise RuntimeError(f"The AI service returned an error (HTTP {error.code}).") from error
    except (URLError, TimeoutError) as error:
        raise RuntimeError("Could not connect to the AI service. Try again shortly.") from error
    except (json.JSONDecodeError, UnicodeDecodeError) as error:
        raise RuntimeError("The AI service returned an unreadable response.") from error

    try:
        content = result["choices"][0]["message"]["content"]
        parsed = json.loads(content)
        corrected = parsed["corrected_text"]
        explanation = parsed.get("explanation", "")
    except (KeyError, IndexError, TypeError, json.JSONDecodeError) as error:
        raise RuntimeError("The AI service returned a response in an unexpected format.") from error
    if not isinstance(corrected, str) or not isinstance(explanation, str):
        raise RuntimeError("The AI service returned a response in an unexpected format.")
    return {"corrected_text": corrected, "explanation": explanation}


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/extension")
def extension_install_page():
    return render_template("extension.html")


@app.get("/download/browser-extension")
def download_browser_extension():
    archive = build_extension_package(firefox=False)
    return send_file(
        archive,
        mimetype="application/zip",
        as_attachment=True,
        download_name="spellingbee-browser-extension.zip",
    )


@app.get("/download/firefox-addon")
def download_firefox_addon():
    archive = build_extension_package(firefox=True)
    return send_file(
        archive,
        mimetype="application/x-xpinstall",
        as_attachment=True,
        download_name="spellingbee-firefox-addon.xpi",
    )


def build_extension_package(firefox: bool) -> io.BytesIO:
    archive = io.BytesIO()
    with ZipFile(archive, "w", ZIP_DEFLATED) as bundle:
        for file_path in sorted(EXTENSION_DIR.iterdir()):
            if not file_path.is_file():
                continue
            if file_path.name == "config.json":
                content = json.dumps({"apiUrl": request.url_root.rstrip("/")}).encode("utf-8")
                bundle.writestr(file_path.name, content)
            elif firefox and file_path.name == "manifest.json":
                manifest = json.loads(file_path.read_text(encoding="utf-8"))
                manifest["browser_specific_settings"] = {
                    "gecko": {
                        "id": "spellingbee-writing-assistant@spellingbee.app",
                        "strict_min_version": "109.0",
                    }
                }
                bundle.writestr(file_path.name, json.dumps(manifest, indent=2))
            else:
                bundle.write(file_path, arcname=file_path.name)
    archive.seek(0)
    return archive


@app.get("/service-worker.js")
def service_worker():
    response = send_from_directory(app.static_folder, "service-worker.js")
    response.headers["Service-Worker-Allowed"] = "/"
    response.headers["Cache-Control"] = "no-cache"
    return response


@app.errorhandler(413)
def request_too_large(_error):
    return jsonify(error="The document is larger than the 10 MB upload limit."), 413


@app.post("/api/proofread")
def proofread():
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or not isinstance(data.get("text"), str):
        return jsonify(error="Enter some text to check."), 400
    text = data["text"].strip()
    if not text:
        return jsonify(error="Enter some text to check."), 400
    if len(text) > 12000:
        return jsonify(error="Text must be 12,000 characters or fewer."), 413
    try:
        return jsonify(proofread_with_ai(text))
    except RuntimeError as error:
        return jsonify(error=str(error)), 503


@app.post("/api/open-document")
def open_document():
    uploaded = request.files.get("document")
    if uploaded is None or not uploaded.filename:
        return jsonify(error="Choose a document to open."), 400
    try:
        text = extract_document_text(uploaded.filename, uploaded.read())
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except ImportError:
        return jsonify(error="Document support is not installed. Reinstall the app requirements."), 503
    except Exception:
        return jsonify(error="This document could not be read. Check that it is a valid, unprotected file."), 400

    text = text.strip()
    if not text:
        return jsonify(error="This document does not contain any text."), 400
    if len(text) > MAX_TEXT_LENGTH:
        return jsonify(error="This document exceeds the 12,000-character limit. Open a shorter document or select an excerpt."), 413
    return jsonify(text=text)


if __name__ == "__main__":
    app.run(host=os.getenv("HOST", "0.0.0.0"), port=int(os.getenv("PORT", "5000")), debug=False)
